"""Super-admin Excel import: read-only preview before any account/database writes."""
import csv
import hashlib
import io
import json
import re
import secrets
import uuid
import zipfile
from datetime import datetime, timezone
from urllib.parse import urlparse

HEADERS = [
    "Participant name", "Participant email", "Organisation", "Password",
    "Confirm password", "Event", "Category", "Title", "Thematic area (optional)",
    "Short summary (optional)", "Full abstract", "Keywords (comma-separated)",
    "Technical help, investment or collaboration sought (optional)",
    "Public poster URL (optional)", "Public video URL (optional)", "Publish immediately",
]
FIELDS = [
    "name", "email", "organisation", "password", "confirm_password", "event",
    "category", "title", "thematic_area", "summary", "abstract", "keywords",
    "support_request", "poster_url", "video_url", "publish",
]
MAX_ROWS = 500
ALIASES = {h.lower(): f for h, f in zip(HEADERS, FIELDS)}
ALIASES.update({"username": "name", "user name": "name", "full name": "name",
                "email": "email", "abstract": "abstract", "poster url": "poster_url",
                "video url": "video_url", "keywords": "keywords", "short summary": "summary",
                "thematic area": "thematic_area", "support request": "support_request"})
SUBMISSION_FIELDS = ["title", "abstract", "thematic_area", "summary", "keywords",
                     "support_request", "poster_url", "video_url"]


def read_excel(data):
    from openpyxl import load_workbook
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Use an .xlsx file smaller than 10 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("The workbook is too large when opened.")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError("Could not read the Excel file. Use the downloaded .xlsx template.") from None
    try:
        sheet = workbook["Import"] if "Import" in workbook.sheetnames else workbook.worksheets[0]
        if (sheet.max_column or 0) > 50 or (sheet.max_row or 0) > 10000:
            raise ValueError("Use at most 50 columns and 500 filled rows.")
        cells = sheet.iter_rows()
        header = next(cells, ())
        if len(header) > 50:
            raise ValueError("Use at most 50 columns.")
        mapping = {}
        for index, cell in enumerate(header):
            label = str(cell.value or "").strip().lower()
            if not label:
                continue
            if label not in ALIASES:
                raise ValueError("Unknown column: " + str(cell.value))
            field = ALIASES[label]
            if field in mapping.values():
                raise ValueError("Duplicate column: " + str(cell.value))
            mapping[index] = field
        if not {"name", "email"}.issubset(mapping.values()):
            raise ValueError("Participant name (or Username) and Participant email (or Email) columns are required.")
        rows = []
        for row_number, cells in enumerate(cells, start=2):
            if row_number > 10000 or len(cells) > 50:
                raise ValueError("Use at most 50 columns and 500 filled rows.")
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            row = {field: "" for field in FIELDS}
            row["_row"] = row_number
            for index, cell in enumerate(cells):
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    raise ValueError(f"Row {row_number}: formulas are not accepted. Paste values only.")
                if index not in mapping:
                    if str(cell.value).strip():
                        raise ValueError(f"Row {row_number}: data appears in a column without a header.")
                    continue
                value = cell.value
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                field = mapping[index]
                row[field] = str(value) if field in ("password", "confirm_password") else str(value).strip()
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError("Import at most 500 filled rows at a time.")
        if not rows:
            raise ValueError("The import sheet has no participant rows.")
        return rows
    finally:
        workbook.close()


def list_accounts(admin):
    users = {}
    page = 1
    while True:
        response = admin.auth.admin.list_users(page=page, per_page=1000)
        batch = response if isinstance(response, list) else response.users
        for user in batch:
            if user.email:
                users[user.email.strip().lower()] = user.id
        if len(batch) < 1000:
            return users
        page += 1


def resolve_event(value, events, default_id):
    if not value:
        return default_id
    matches = [event["id"] for event in events
               if event["id"] == value or event["title"].strip().casefold() == value.casefold()]
    if len(matches) != 1:
        raise ValueError("Event must match one event title or ID.")
    return matches[0]


def plan_import(rows, accounts, events, default_event=None):
    plans, errors = [], []
    first_accounts = set(accounts)
    passwords = {}
    for row in rows:
        email = row["email"].lower()
        if email not in accounts and row["password"]:
            if email in passwords and passwords[email] != row["password"]:
                errors.append({"Row": row["_row"], "Error": "Different passwords supplied for the same new email."})
            passwords[email] = row["password"]
    limits = {"name": 120, "email": 254, "organisation": 160, "title": 200,
              "thematic_area": 160, "summary": 300, "abstract": 8000,
              "keywords": 300, "support_request": 2000}
    for source in rows:
        row = dict(source)
        row["email"] = row["email"].lower()
        issues = []
        if len(row["name"]) < 2:
            issues.append("Participant name / Username is required (at least 2 characters).")
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", row["email"]):
            issues.append("Enter a valid participant email.")
        for field, limit in limits.items():
            if len(row[field]) > limit:
                issues.append(f"{field.replace('_', ' ').title()} exceeds {limit} characters.")
        if row["email"] not in accounts:
            password = passwords.get(row["email"], "")
            if password and len(password) < 8:
                issues.append("Passwords must have at least 8 characters.")
            if row["confirm_password"] and row["confirm_password"] != password:
                issues.append("Confirm password does not match.")
        else:
            password = ""
        has_submission = any(row[field] for field in SUBMISSION_FIELDS)
        event_id = None
        category = row["category"].strip().lower().replace("_", " ") or "project idea"
        if category not in ("project idea", "oral presentation", "poster"):
            issues.append("Category must be Project idea, Oral presentation or Poster.")
        publish_value = row["publish"].lower()
        if publish_value not in ("", "true", "false", "yes", "no", "1", "0"):
            issues.append("Publish immediately must be Yes or No, or blank.")
        if has_submission:
            try:
                event_id = resolve_event(row["event"], events, default_event)
                if not event_id or not any(e["id"] == event_id for e in events):
                    issues.append("Choose a default event or fill the Event cell.")
            except ValueError as exc:
                issues.append(str(exc))
            if row["title"] and len(row["title"]) < 5:
                issues.append("A supplied title must have at least 5 characters.")
            if row["abstract"] and len(row["abstract"]) < 30:
                issues.append("A supplied abstract must have at least 30 characters.")
            for field in ("poster_url", "video_url"):
                if row[field]:
                    parsed = urlparse(row[field])
                    if parsed.scheme != "https" or not parsed.hostname:
                        issues.append("Poster and video links must be public HTTPS URLs.")
        if issues:
            errors.append({"Row": row["_row"], "Error": " ".join(issues)})
            continue
        incomplete = has_submission and (not row["title"] or not row["abstract"])
        plans.append({
            "row": row["_row"], "name": row["name"], "email": row["email"],
            "organisation": row["organisation"], "password": password,
            "new_account": row["email"] not in first_accounts,
            "submission": has_submission, "incomplete": incomplete,
            "event_id": event_id, "category": category.replace(" ", "_"),
            "title": row["title"] or f"Contribution by {row['name']}",
            "abstract": row["abstract"] or "Abstract not provided. The author needs to add their abstract.",
            "summary": row["summary"], "thematic_area": row["thematic_area"],
            "keywords": [word.strip() for word in row["keywords"].split(",") if word.strip()],
            "support_request": row["support_request"], "poster_url": row["poster_url"],
            "video_url": row["video_url"],
            "status": "approved" if publish_value in ("yes", "true", "1") and not incomplete else "submitted",
        })
        first_accounts.add(row["email"])
    return plans, errors


def preview_rows(plans):
    return [{"Row": p["row"], "Participant name": p["name"], "Email": p["email"],
             "Account": "Create" if p["new_account"] else "Use existing",
             "Password": ("Provided" if p["password"] else "Generate") if p["new_account"] else "Keep existing",
             "Submission": (("Incomplete — review only" if p["incomplete"] else p["status"])
                            if p["submission"] else "No submission"),
             "Title": p["title"] if p["submission"] else ""} for p in plans]


def submission_payload(plan, owner_id, reviewer_id):
    fields = {
        "owner_id": owner_id, "event_id": plan["event_id"], "title": plan["title"],
        "abstract_text": plan["abstract"], "short_summary": plan["summary"] or None,
        "thematic_area": plan["thematic_area"] or None, "keywords": plan["keywords"],
        "poster_url": plan["poster_url"] or None, "video_url": plan["video_url"] or None,
        "presentation_category": plan["category"], "support_request": plan["support_request"] or None,
    }
    digest = hashlib.sha256(json.dumps(fields, sort_keys=True).encode()).hexdigest()
    # Same content gets the same primary key on retries and concurrent imports.
    submission_id = str(uuid.uuid5(uuid.NAMESPACE_URL, "octa-bulk-submission:" + digest))
    now = datetime.now(timezone.utc).isoformat()
    return {**fields, "id": submission_id, "slug": "bulk-" + submission_id.replace("-", ""),
            "status": plan["status"], "submitted_at": now,
            "approved_by": reviewer_id if plan["status"] == "approved" else None,
            "approved_at": now if plan["status"] == "approved" else None}


def import_rows(admin, plans, accounts, reviewer_id, credentials=None, progress=None):
    results = []
    credentials = credentials if credentials is not None else []
    for index, plan in enumerate(plans):
        result = {"Row": plan["row"], "Email": plan["email"], "Account": "",
                  "Submission": "Not requested", "Error": ""}
        stage = "account"
        try:
            owner_id = accounts.get(plan["email"])
            if not owner_id:
                password = plan["password"] or secrets.token_urlsafe(18)
                response = admin.auth.admin.create_user({
                    "email": plan["email"], "password": password, "email_confirm": True,
                    "user_metadata": {"full_name": plan["name"]},
                })
                if not response.user:
                    raise ValueError("Account was not returned.")
                owner_id = response.user.id
                accounts[plan["email"]] = owner_id
                result["Account"] = "Created"
                # Preserve credentials immediately, even if the next database write fails.
                credentials.append({"Participant name": plan["name"], "Email": plan["email"],
                                    "Password": password})
            else:
                result["Account"] = "Existing — password unchanged"
            stage = "profile"
            existing = admin.table("profiles").select("id").eq("id", owner_id).execute().data or []
            if not existing or result["Account"] == "Created":
                admin.table("profiles").upsert({
                    "id": owner_id, "full_name": plan["name"],
                    "organisation": plan["organisation"],
                }, on_conflict="id").execute()
            types = admin.table("platform_user_types").select("user_id").eq(
                "user_id", owner_id).execute().data or []
            if not types:
                admin.table("platform_user_types").insert({
                    "user_id": owner_id, "user_type": "project_owner",
                }).execute()
            if plan["submission"]:
                stage = "submission"
                payload = submission_payload(plan, owner_id, reviewer_id)
                existing = admin.table("submissions").select("id").eq("id", payload["id"]).execute().data or []
                if existing:
                    result["Submission"] = "Skipped — already imported"
                else:
                    try:
                        admin.table("submissions").insert(payload).execute()
                        result["Submission"] = "Published" if plan["status"] == "approved" else "Added for review"
                    except Exception:
                        # Handles an uncertain response or another simultaneous import.
                        found = admin.table("submissions").select("id").eq("id", payload["id"]).execute().data or []
                        if not found:
                            raise
                        result["Submission"] = "Skipped — already imported"
        except Exception as exc:
            # Do not echo API exception messages, which could contain passwords.
            result["Error"] = f"{stage.title()} failed ({type(exc).__name__}). Successful earlier steps remain saved. Retry this row."
        results.append(result)
        if progress:
            progress((index + 1) / len(plans))
    return results, credentials


def credentials_csv(rows):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=["Participant name", "Email", "Password"])
    writer.writeheader()
    for row in rows:
        # Avoid spreadsheet formula interpretation in a credentials CSV.
        writer.writerow({key: "'" + str(value) if str(value).startswith(("=", "+", "-", "@"))
                         else value for key, value in row.items()})
    return stream.getvalue().encode("utf-8-sig")
